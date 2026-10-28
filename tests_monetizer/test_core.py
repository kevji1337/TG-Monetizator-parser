import io
import json
import tempfile
import threading
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import MagicMock, patch

from openpyxl import load_workbook

from monetizer.audit import (AuditError, Cancelled, WebAuditor, normalize_username,
                             parse_channel_html, parse_count, retry_after_seconds)
from monetizer.export import COLUMNS, csv_bytes, xlsx_bytes
from monetizer.keywords import expand_keywords
from monetizer.pitches import generate_pitch
from monetizer.runner import JobRunner
from monetizer.settings import Settings

NOW = datetime(2026, 10, 15, 12, tzinfo=timezone.utc)


def preview(subs="4K", views=700, days=None, about=None):
    days = days if days is not None else [1, 2, 3, 4, 5, 6]
    about = about if about is not None else ('Автор <a href="https://t.me/expert_contact">@expert_contact</a>'
                                            '<a href="https://youtube.com/shorts/abc">YouTube</a>')
    posts = "".join(f'<div class="tgme_widget_message extra" data-post="expert_blog/{i}">'
                    f'<a class="tgme_widget_message_date"><time datetime="{(NOW - timedelta(days=day)).isoformat()}"></time></a>'
                    f'<span class="tgme_widget_message_views">{views}</span>'
                    '<div class="tgme_widget_message_text">Полезный пост</div></div>' for i, day in enumerate(days))
    return ('<div class="tgme_channel_info"><div class="tgme_channel_info_header_title">Эксперт &amp; практика</div>'
            f'<div class="tgme_channel_info_counter"><span class="counter_value">{subs}</span>'
            '<span class="counter_type">subscribers</span></div>'
            f'<div class="tgme_channel_info_description">{about}</div></div>{posts}')


class AuditTests(unittest.TestCase):
    def test_counters_and_usernames(self):
        for text, expected in [("1,234", 1234), ("3\u202f500", 3500), ("1.2K", 1200),
                               ("1,2M", 1200000), ("2 тыс.", 2000), ("—", None)]:
            self.assertEqual(parse_count(text), expected)
        self.assertEqual(normalize_username("https://t.me/s/Expert_blog/15"), "expert_blog")
        for invalid in ("../../etc", "https://evil.test/abcd", "@a", "https://t.me/+invitation"):
            with self.assertRaises(ValueError):
                normalize_username(invalid)

    def test_live_channel_and_contact(self):
        row = parse_channel_html(preview(), "expert_blog", now=NOW)
        self.assertEqual(row["status"], "top")
        self.assertEqual(row["title"], "Эксперт & практика")
        self.assertEqual(row["avg_views"], 700)
        self.assertEqual(row["median_views"], 700)
        self.assertEqual(row["avg_gap_days"], 1)
        self.assertEqual(row["last_post"], (NOW - timedelta(days=1)).isoformat())
        self.assertEqual(row["author_contact"], "@expert_contact")

    def test_subscriber_threshold_and_no_rounding_promotion(self):
        small = parse_channel_html(preview(subs="4,999", views=499), "expert_blog", now=NOW)
        self.assertEqual(small["status"], "rejected")
        self.assertEqual(small["norm"], 10)
        row = parse_channel_html(preview(subs="5K", views=500), "expert_blog", now=NOW)
        self.assertEqual(row["status"], "top")
        self.assertEqual(row["norm"], 10)

    def test_age_cadence_and_small_sample(self):
        for days, status in [([15, 16, 17, 18, 19, 20], "rejected"),
                             ([8, 9, 10, 11, 12, 13], "review"),
                             ([1, 7, 13, 19, 25, 31], "review"),
                             ([.01, .1, .2, .3, .4, .5], "review")]:
            self.assertEqual(parse_channel_html(preview(days=days), "expert_blog", now=NOW)["status"], status)

    def test_missing_metrics_are_not_zero_or_fraud(self):
        html = preview(subs="unknown").replace('class="tgme_widget_message_views"', 'class="missing"')
        row = parse_channel_html(html, "expert_blog", now=NOW)
        self.assertEqual(row["status"], "review")
        self.assertIsNone(row["subs"])
        self.assertIsNone(row["avg_views"])
        row = parse_channel_html("<html>Access denied</html>", "expert_blog", now=NOW)
        self.assertEqual(row["status"], "error")

    def test_link_domains_bot_contact_and_self_are_filtered(self):
        row = parse_channel_html(preview(about='@expert_blog @fakebot <a href="https://youtube.com.evil.test/watch">Link</a>'),
                                 "expert_blog", now=NOW)
        self.assertEqual(row["status"], "recommended")
        self.assertFalse(row["youtube"])
        self.assertFalse(row["author_contact"])

    def test_duplicate_message_and_fresh_views(self):
        html = preview(days=[0, 1, 2, 3, 4, 5, 6])
        html = html.replace('data-post="expert_blog/0"', 'data-post="expert_blog/1"')
        row = parse_channel_html(html, "expert_blog", now=NOW)
        self.assertEqual(row["post_count"], 6)
        self.assertEqual(row["eligible_posts"], 6)

    def test_keywords_and_pitch_only_use_available_evidence(self):
        queries = expand_keywords("дизайн", limit=12)
        self.assertTrue(any("framer" in q.casefold() for q in queries))
        self.assertEqual(len(queries), len(set(queries)))
        self.assertEqual(expand_keywords("гончарное дело", limit=10), ["гончарное дело"])
        row = parse_channel_html(preview(about=""), "expert_blog", now=NOW)
        pitch = generate_pitch(row, "дизайн", sender="Анна")
        self.assertIn("Анна", pitch)
        self.assertNotIn("YouTube", pitch)
        self.assertIn("экспертную практику", pitch)

    def test_exports_preserve_strings_and_csv_bom(self):
        row = parse_channel_html(preview(), "expert_blog", now=NOW)
        row["title"] = '=HYPERLINK("evil")'
        data = csv_bytes([row])
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))
        self.assertIn("'=HYPERLINK", data.decode("utf-8-sig"))
        workbook = load_workbook(io.BytesIO(xlsx_bytes([row])))
        cell = workbook.active.cell(2, list(COLUMNS).index("title") + 1)
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(workbook.active.max_row, 2)


class TransportTests(unittest.TestCase):
    def response(self, status=200, html=None, retry="0"):
        response = MagicMock()
        response.__enter__.return_value = response
        response.status_code = status
        response.url = "https://t.me/s/expert_blog"
        response.headers = {"Retry-After": retry}
        response.iter_content.return_value = [(html or preview()).encode()]
        return response

    def test_retries_429_and_no_retry_403(self):
        auditor = WebAuditor(Settings())
        session = MagicMock()
        session.get.side_effect = [self.response(429), self.response(200)]
        with patch.object(auditor, "_session", return_value=session), patch.object(auditor.limiter, "acquire"):
            self.assertIn("tgme_channel_info", auditor.fetch("expert_blog"))
        self.assertEqual(session.get.call_count, 2)
        session.get.side_effect = [self.response(403)]
        session.get.reset_mock()
        with patch.object(auditor, "_session", return_value=session), patch.object(auditor.limiter, "acquire"):
            self.assertEqual(auditor.audit("expert_blog")["error_code"], "http_403")
        self.assertEqual(session.get.call_count, 1)

    def test_long_retry_after_blocks_batch_without_long_wait(self):
        auditor, session = WebAuditor(Settings()), MagicMock()
        session.get.return_value = self.response(429, retry="3600")
        with patch.object(auditor, "_session", return_value=session), patch.object(auditor.limiter, "acquire"):
            self.assertEqual(auditor.audit("expert_blog")["error_code"], "http_429")
            self.assertEqual(auditor.audit("other_blog")["error_code"], "http_429")
        self.assertEqual(session.get.call_count, 1)
        self.assertEqual(retry_after_seconds("0"), 0)

    def test_cancel_before_network(self):
        stop = threading.Event()
        stop.set()
        with self.assertRaises(Cancelled):
            WebAuditor(stop=stop).fetch("expert_blog")


class RunnerTests(unittest.TestCase):
    def setUp(self):
        self.lock_directory = tempfile.TemporaryDirectory()
        self.lock_patch = patch("monetizer.runner.RESULTS_DIR", Path(self.lock_directory.name))
        self.lock_patch.start()

    def tearDown(self):
        self.lock_patch.stop()
        self.lock_directory.cleanup()

    def test_worker_failure_keeps_other_rows_and_restores_checkpoint(self):
        with tempfile.TemporaryDirectory() as directory:
            checkpoint = Path(directory) / "checkpoint.json"
            runner = JobRunner(checkpoint)
            def audit(username):
                if username == "broken_blog":
                    raise ValueError("Broken worker")
                return parse_channel_html(preview(), username, now=NOW)
            with patch("monetizer.runner.WebAuditor") as factory, patch("monetizer.runner.save_reports", return_value=(["test.csv"], [])):
                factory.return_value.audit.side_effect = audit
                runner.start(Settings(target_count=2), niche="Дизайн", usernames=["expert_blog", "broken_blog"])
                runner.thread.join(10)
            self.assertFalse(runner.thread.is_alive())
            state = runner.snapshot()
            self.assertEqual(state["status"], "completed")
            self.assertEqual(len(state["rows"]), 2)
            self.assertEqual(sum(r["status"] == "error" for r in state["rows"]), 1)
            self.assertEqual(JobRunner(checkpoint).snapshot()["rows"], state["rows"])

    def test_stop_and_resume_skips_finished_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            runner = JobRunner(Path(directory) / "checkpoint.json")
            first = threading.Event()
            def audit(username):
                if username == "expert_blog":
                    first.set()
                    return parse_channel_html(preview(), username, now=NOW)
                runner.stop_event.wait(2)
                raise Cancelled()
            with patch("monetizer.runner.WebAuditor") as factory, patch("monetizer.runner.save_reports", return_value=([], [])):
                factory.return_value.audit.side_effect = audit
                runner.start(Settings(target_count=2, workers=1), niche="Дизайн", usernames=["expert_blog", "other_blog"])
                self.assertTrue(first.wait(2))
                runner.stop()
                runner.thread.join(10)
                self.assertEqual(runner.snapshot()["status"], "cancelled")
                factory.return_value.audit.side_effect = lambda u: {**parse_channel_html(preview(), "expert_blog", now=NOW), "username": "@" + u}
                runner.start(resume=True)
                runner.thread.join(10)
            self.assertEqual(runner.snapshot()["status"], "completed")
            self.assertEqual(len(runner.snapshot()["rows"]), 2)


if __name__ == "__main__":
    unittest.main()
