import csv
import io
import json
import re
from datetime import datetime
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
RESULTS_DIR = PROJECT_ROOT / "results" / "monetizer"
REPORTS_DIR = PROJECT_ROOT / 'reports'
COLUMNS = {
    "verdict": "Вердикт", "username": "Канал", "title": "Название", "subs": "Подписчики",
    "avg_views": "Средние просмотры поста", "median_views": "Медиана просмотров",
    "view_ratio": "Просмотры/подписчики, %", "norm": "Порог, %", "freq": "Частота",
    "last_post": "Последний пост (UTC)", "days_since_last": "Дней без постов",
    "post_count": "Постов в выборке", "eligible_posts": "Постов с просмотрами",
    "youtube": "YouTube", "instagram": "Instagram", "author_contact": "Кандидат контакта",
    "contact_candidates": "Все кандидаты", "contact_status": "Статус контакта",
    "producer_status": "Наличие продюсера", "sales_links": "Ссылки продаж",
    "link": "Ссылка канала", "source_queries": "Поисковые запросы", "reasons": "Основания",
    "error_code": "Код ошибки", "error": "Ошибка", "pitch": "Черновик сообщения",
    "audited_at": "Проверено (UTC)", "status": "Код статуса", "external_links": "Все внешние ссылки",
    "pitch_scenario": "Сценарий сообщения", "pitch_source": "Источник наблюдения",
    "metric_status": "Результат проверки метрик", "channel_type": "Тип канала", "relevance": "Соответствие теме",
    "classification": "Основания классификации", "offers_summary": "Предложения автора", "offers": "Предложения с источниками",
    "work_status_label": "Рабочий статус", "notes": "Заметки", "ai_review": "Отдельная оценка AI",
    "source_categories": "Тематики TGStat",
    'tgstat_filters': 'Фильтры поиска TGStat',
    'assessment_stale': 'Автооценка устарела',
    'ai_audit_status': 'Статус AI-аудита API', 'ai_detail': 'Причина пропуска или ошибки AI',
    'api_ai_review': 'Разбор и черновик Gemini API', 'local_verdict': 'Вердикт локальных правил',
    'ai_queued': 'Ожидает AI-разбора',
    'ai_priority_score': 'Приоритет очереди AI, 0–100', 'ai_priority_reason': 'Причины приоритета AI',
    'selection_bucket_label': 'Главная причина результата',
    'next_contact_date': 'Дата следующего контакта', 'next_action_label': 'Следующее действие',
    'contact_plan_note': 'Пояснение к действию', 'contact_plan_state': 'Срок действия',
    'last_action_at': 'Последнее выполненное действие (UTC)',
    'activity_status': 'Проверка активности', 'activity_ratio': 'Просмотры выбранного поста / подписчики, %',
    'activity_post_age_hours': 'Возраст поста при замере, ч', 'activity_post_url': 'Пост для оценки активности',
    'activity_views': 'Просмотры выбранного поста', 'activity_subscribers': 'Подписчики при замере',
    'activity_observed_at': 'Время замера активности (UTC)',
    'activity_target_age_hours': 'Целевой возраст поста, ч', 'activity_age_tolerance_hours': 'Отклонение возраста поста, ч',
    'activity_evidence': 'Замер активности поста',
}


def _value(value):
    return json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value


def csv_bytes(rows):
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter=";")
    writer.writerow(COLUMNS.values())
    for row in rows:
        values = []
        for key in COLUMNS:
            value = _value(row.get(key, ""))
            if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
                value = "'" + value
            values.append(value)
        writer.writerow(values)
    return output.getvalue().encode("utf-8-sig")


def xlsx_bytes(rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Каналы"
    sheet.append(list(COLUMNS.values()))
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="153A45")
    colors = {"top": "DBF0DD", "recommended": "EDF7ED", "review": "FFF2CC", "rejected": "FBE4E4", "error": "EEEEEE"}
    for row in rows:
        sheet.append([_value(row.get(key)) for key in COLUMNS])
        for column, key in enumerate(COLUMNS, 1):
            cell = sheet.cell(sheet.max_row, column)
            if isinstance(cell.value, str):
                # Store external strings literally, including leading @ or =.
                cell.data_type = "s"
            cell.fill = PatternFill("solid", fgColor=colors.get(row.get("status"), "FFFFFF"))
            if key in ("link", "youtube", "instagram") and str(cell.value or "").startswith("https://"):
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, key in enumerate(COLUMNS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = 55 if key in ("pitch", "reasons", "title") else 24
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def atomic_json(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def report_stem(niche):
    name = re.sub(r"[^\wа-яА-ЯёЁ-]+", "_", niche, flags=re.UNICODE).strip("_")[:65] or "каналы"
    return f"Монетизация_{name}_{datetime.now().strftime('%Y%m%d_%H%M%S_%f')}"


def save_reports(rows, niche):
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    stem = report_stem(niche)
    payloads = {"csv": csv_bytes(rows), "xlsx": xlsx_bytes(rows)}
    paths, warnings = [], []
    for extension, data in payloads.items():
        path = REPORTS_DIR / f"{stem}.{extension}"
        try:
            path.write_bytes(data)
            paths.append(str(path))
        except OSError as exc:
            warnings.append(f"Не удалось сохранить {path}: {exc}")
    if not paths:
        raise OSError("; ".join(warnings))
    return paths, warnings
